"""Bảng tra mã quốc tịch và địa danh hành chính Việt Nam."""

import streamlit as st
import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from copy import copy
import xlrd, datetime, io, zipfile, base64, os, json
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.lib.colors import white, black
import unicodedata as _ud, re as _re

from .common import _strip_accents
from .assets import load_template

# ── Lookup tables ─────────────────────────────────────────────────────────
# Full nationality mapping (normalized keys → "CODE - Name") — 350 entries
import unicodedata as _ud, re as _re



def _norm_nat_legacy(s):
    """Cách chuẩn hoá CŨ (nuốt mất chữ đ). Các khoá trong NAT_NORM bên dưới
    được gõ theo đúng cách này (VD "Đài Loan" → "ailoan"), nên vẫn phải tra
    được bằng nó — xem lookup_nat_kbtt()."""
    s = str(s).lower().strip()
    s = _ud.normalize('NFD', s)
    s = ''.join(c for c in s if _ud.category(c) != 'Mn')
    return _re.sub(r'[^a-z0-9]', '', s)

def _norm_nat(s):
    return _re.sub(r'[^a-z0-9]', '', _strip_accents(s).lower().strip())

NAT_NORM = {
    "achentina": "ARG - Argentina",
    "acmenia": "ARM - Armenia",
    "adecbaigian": "AZE - Azerbaijan",
    "aicap": "EGY - Egypt",
    "airolen": "IRL - Ireland",
    "aixolen": "ISL - Iceland",
    "albania": "ALB - Albania",
    "anbani": "ALB - Albania",
    "andorra": "AND - Andorra",
    "anggola": "AGO - Angola",
    "angola": "AGO - Angola",
    "anh": "GBR - United Kingdom",
    "anmach": "DNK - Denmark",
    "ano": "IND - India",
    "ao": "AUT - Austria",
    "aosip": "CYP - Cyprus",
    "arapthongnhat": "ARE - United Arab Emirates",
    "arapxaui": "SAU - Saudi Arabia",
    "argentina": "ARG - Argentina",
    "armenia": "ARM - Armenia",
    "australia": "AUS - Australia",
    "austria": "AUT - Austria",
    "azerbaijan": "AZE - Azerbaijan",
    "bacbaot": "BRB - Barbados",
    "bahama": "BHS - Bahamas",
    "bahamas": "BHS - Bahamas",
    "bahrain": "BHR - Bahrain",
    "balan": "POL - Poland",
    "bangladesh": "BGD - Bangladesh",
    "banglaet": "BGD - Bangladesh",
    "barain": "BHR - Bahrain",
    "barbados": "BRB - Barbados",
    "becmua": "BMU - Bermuda",
    "belarus": "BLR - Belarus",
    "belarut": "BLR - Belarus",
    "belgium": "BEL - Belgium",
    "belixe": "BLZ - Belize",
    "belize": "BLZ - Belize",
    "benanh": "BEN - Benin",
    "benin": "BEN - Benin",
    "bermuda": "BMU - Bermuda",
    "bhutan": "BTN - Bhutan",
    "bi": "BEL - Belgium",
    "boaonha": "PRT - Portugal",
    "bolivia": "BOL - Bolivia",
    "bosniaandherzegovina": "BIH - Bosnia and Herzegovina",
    "botswana": "BWA - Botswana",
    "botxoana": "BWA - Botswana",
    "boxniahecdegovina": "BIH - Bosnia and Herzegovina",
    "bradin": "BRA - Brazil",
    "brazil": "BRA - Brazil",
    "britishindiaoceanterritory": "IOT - British India Ocean Territory",
    "bulgaria": "BGR - Bulgaria",
    "bungari": "BGR - Bulgaria",
    "buockinaphaxo": "BFA - Burkina Faso",
    "burkinafaso": "BFA - Burkina Faso",
    "burundi": "BDI - Burundi",
    "buruni": "BDI - Burundi",
    "butan": "BTN - Bhutan",
    "cameroon": "CMR - Cameroon",
    "camorun": "CMR - Cameroon",
    "canada": "CAN - Canada",
    "capeverde": "CPV - Cape Verde",
    "capve": "CPV - Cape Verde",
    "chad": "TCD - Chad",
    "charapxyri": "SYR - Syrian Arab Republic",
    "chdcndtrieutien": "PRK - Korea Democratic Peoples Republic of",
    "chhanquoc": "KOR - Korea (South)",
    "chhoigiaoiran": "IRN - Iran Ilasmic Republic of",
    "chile": "CHL - Chile",
    "china": "CHN - China",
    "chinataiwan": "CHN - China",
    "chlienbanguc": "D - Germany",
    "chmaxeonia": "MKD - Macedonia",
    "chominicana": "DMA - Dominica",
    "colombia": "COL - Colombia",
    "como": "COM - Comoros",
    "comoros": "COM - Comoros",
    "conggo": "COG - Congo",
    "conghoasec": "CZE - Czech Republic",
    "congo": "COG - Congo",
    "congquocanora": "AND - Andorra",
    "congquoclichtenxten": "LIE - Liechtenstein",
    "congquocmonaco": "MCO - Monaco",
    "cooet": "KWT - Kuwait",
    "costarica": "CRI - Costa Rica",
    "cotedivoire": "CIV - Cote d' Ivoire",
    "cotivoa": "CIV - Cote d' Ivoire",
    "coxtarica": "CRI - Costa Rica",
    "croatia": "HRV - Croatia",
    "cuba": "CUB - Cuba",
    "cyprus": "CYP - Cyprus",
    "czechrepublic": "CZE - Czech Republic",
    "dambia": "ZMB - Zambia",
    "denmark": "DNK - Denmark",
    "dimbabue": "ZWE - Zimbabwe",
    "djibouti": "DJI - Djibouti",
    "dominica": "DMA - Dominica",
    "dominicana": "DMA - Dominica",
    "ecuador": "ECU - Ecuador",
    "ecuao": "ECU - Ecuador",
    "egypt": "EGY - Egypt",
    "elsalvador": "SLV - El Salvado",
    "enxanvao": "SLV - El Salvado",
    "equatorialguinea": "GNQ - Equatorial Guinea",
    "eritoria": "ERI - Eritrea",
    "eritrea": "ERI - Eritrea",
    "estonia": "EST - Estonia",
    "ethiopia": "ETH - Ethiopia",
    "etiopia": "ETH - Ethiopia",
    "extonia": "EST - Estonia",
    "fiji": "FJI - Fiji",
    "finland": "FIN - Finland",
    "france": "FRA - France",
    "francemetropolitan": "FRA - France",
    "gabon": "GAB - Gabon",
    "gabong": "GAB - Gabon",
    "gambia": "GMB - Gambia",
    "gana": "GHA - Ghana",
    "georgia": "GEO - Georgia",
    "germany": "D - Germany",
    "ghana": "GHA - Ghana",
    "ghine": "GIN - Guinea",
    "ghinebitxao": "GNB - Guinea-Bissau",
    "ghinexichao": "GNQ - Equatorial Guinea",
    "giamahiriiaaraplibinhandan": "LBY - Libyan Arab Jamahiriya",
    "gibraltar": "GIB - Gibraltar",
    "gibranta": "GIB - Gibraltar",
    "goatemala": "GTM - Guatemala",
    "greece": "GRC - Greece",
    "greenland": "GRL - Greenland",
    "grenaa": "GRD - Grenada",
    "grenada": "GRD - Grenada",
    "grinlon": "GRL - Greenland",
    "grudia": "GEO - Georgia",
    "guatemala": "GTM - Guatemala",
    "guina": "GUY - Guyana",
    "guinea": "GIN - Guinea",
    "guineabissau": "GNB - Guinea-Bissau",
    "guyana": "GUY - Guyana",
    "haiti": "HTI - Haiti",
    "halan": "NLD - Netherland",
    "hanquoc": "KOR - Korea (South)",
    "honduras": "HND - Honduras",
    "hondurat": "HND - Honduras",
    "hylap": "GRC - Greece",
    "iaphanthuoclienhiepanh": "GBD - United Kingdom British Territories Citizen",
    "ibouti": "DJI - Djibouti",
    "iceland": "ISL - Iceland",
    "india": "IND - India",
    "indonesia": "IDN - Indonesia",
    "inonexia": "IDN - Indonesia",
    "irac": "IRQ - Iraq",
    "iran": "IRN - Iran Ilasmic Republic of",
    "iraq": "IRQ - Iraq",
    "ireland": "IRL - Ireland",
    "israel": "ISR - Israel",
    "italia": "ITA - Italy",
    "italy": "ITA - Italy",
    "ixraen": "ISR - Israel",
    "jamaica": "JAM - Jamaica",
    "japan": "JPN - Japan",
    "jocan": "JOR - Jordan",
    "jordan": "JOR - Jordan",
    "kadacxtan": "KAZ - Kazakhstan",
    "kazakhstan": "KAZ - Kazakhstan",
    "kenia": "KEN - Kenya",
    "kenya": "KEN - Kenya",
    "kiecghidia": "KGZ - Kyrgyzstan",
    "kiribati": "KIR - Kiribati",
    "koreademocraticpeoplesrepublic": "PRK - Korea Democratic Peoples Republic of",
    "koreasouth": "KOR - Korea (South)",
    "kosovo": "RKS - Kosovo",
    "kuwait": "KWT - Kuwait",
    "kyrgyzstan": "KGZ - Kyrgyzstan",
    "latvia": "LVA - Latvia",
    "lebanon": "LBN - Lebanon",
    "lesotho": "LSO - Lesotho",
    "lexotho": "LSO - Lesotho",
    "liban": "LBN - Lebanon",
    "liberia": "LBR - Liberia",
    "libya": "LBY - Libyan Arab Jamahiriya",
    "liechtenstein": "LIE - Liechtenstein",
    "lienbangnga": "RUS - Russia",
    "lithuania": "LTU - Lithuania",
    "luxembourg": "LUX - Luxembourg",
    "luychxembua": "LUX - Luxembourg",
    "maagaxca": "MDG - Madagascar",
    "macedonia": "MKD - Macedonia",
    "madagascar": "MDG - Madagascar",
    "malaixia": "MYS - Malaysia",
    "malauy": "MWI - Malawi",
    "malawi": "MWI - Malawi",
    "malaysia": "MYS - Malaysia",
    "maldives": "MDV - Maldives",
    "mali": "MLI - Mali",
    "malta": "MLT - Malta",
    "manivo": "MDV - Maldives",
    "manta": "MLT - Malta",
    "maroc": "MAR - Morocco",
    "marshallislands": "MHL - Marshall Islands",
    "mauritania": "MRT - Mauritania",
    "mauritius": "MUS - Mauritius",
    "mexico": "MEX - Mexico",
    "mianma": "MMR - Myanmar",
    "micronesia": "FSM - Micronesia",
    "modambich": "MOZ - Mozambique",
    "moldova": "MDA - Moldova",
    "monaco": "MCO - Monaco",
    "mongco": "MNG - Mongolia",
    "mongolia": "MNG - Mongolia",
    "monova": "MDA - Moldova",
    "montenegro": "MNE - Montenegro",
    "montserrat": "MSR - Montserrat",
    "monxerat": "MSR - Montserrat",
    "moratani": "MRT - Mauritania",
    "morixo": "MUS - Mauritius",
    "morocco": "MAR - Morocco",
    "mozambique": "MOZ - Mozambique",
    "my": "USA - United States of America",
    "myanmarburma": "MMR - Myanmar",
    "namibia": "NAM - Namibia",
    "nauru": "NRU - Nauru",
    "nauy": "NOR - Norway",
    "nepal": "NPL - Nepal",
    "nepan": "NPL - Nepal",
    "netherland": "NLD - Netherland",
    "netherlandantilles": "NLD - Netherland",
    "newzealand": "NZL - New Zealand",
    "nhatban": "JPN - Japan",
    "nicaragoa": "NIC - Nicaragua",
    "nicaragua": "NIC - Nicaragua",
    "niger": "NER - Niger",
    "nigeria": "NGA - Nigeria",
    "nigie": "NER - Niger",
    "nigieria": "NGA - Nigeria",
    "niudilan": "NZL - New Zealand",
    "norway": "NOR - Norway",
    "oman": "OMN - Oman",
    "ominica": "DMA - Dominica",
    "ongtimo": "TLS - Timor Leste",
    "oxtraylia": "AUS - Australia",
    "pakistan": "PAK - Pakistan",
    "pakixtan": "PAK - Pakistan",
    "palau": "PLW - Palau",
    "palestine": "PSE - Palestine",
    "palextin": "PSE - Palestine",
    "panama": "PAN - Panama",
    "papuanewguinea": "PNG - Papua New Guinea",
    "papuaniughine": "PNG - Papua New Guinea",
    "paragoay": "PRY - Paraguay",
    "paraguay": "PRY - Paraguay",
    "peru": "PER - Peru",
    "phanlan": "FIN - Finland",
    "phap": "FRA - France",
    "philippin": "PHL - Philippines",
    "philippine": "PHL - Philippines",
    "poland": "POL - Poland",
    "portugal": "PRT - Portugal",
    "qatar": "QAT - Qatar",
    "quanaoantithuochalan": "NLD - Netherland",
    "quanaomacsan": "MHL - Marshall Islands",
    "quanaonamgrudiavanamsanuych": "GEO - Georgia",
    "quanaoxaysen": "SYC - Seychelles",
    "quata": "QAT - Qatar",
    "romania": "ROU - Romania",
    "ruana": "RWA - Rwanda",
    "rumani": "ROU - Romania",
    "russia": "RUS - Russia",
    "rwanda": "RWA - Rwanda",
    "saintlucia": "LCA - Saint Lucia",
    "sanmarino": "SMR - San Marino",
    "sat": "TCD - Chad",
    "saudiarabia": "SAU - Saudi Arabia",
    "scotland": "SC- - Scotland",
    "senegal": "SEN - Senegal",
    "serbia": "SRB - Serbia",
    "seychelles": "SYC - Seychelles",
    "singapore": "SGP - Singapore",
    "slovakia": "SVK - Slovakia",
    "slovenia": "SVN - Slovenia",
    "somalia": "SOM - Somalia",
    "southgeorgiaandthesouths": "GEO - Georgia",
    "spain": "ESP - Spain",
    "srilanka": "LKA - Sri Lanka",
    "sudan": "SDN - Sudan",
    "suriname": "SUR - Suriname",
    "swaziland": "SWZ - Swaziland",
    "sweden": "SWE - Sweden",
    "switzerland": "CHE - Switzerland",
    "syria": "SYR - Syrian Arab Republic",
    "tagikixtan": "TJK - Tajikistan",
    "tajikistan": "TJK - Tajikistan",
    "taybannha": "ESP - Spain",
    "thailan": "THA - Thailand",
    "thailand": "THA - Thailand",
    "thonhiky": "TUR - Turkey",
    "thuyien": "SWE - Sweden",
    "thuysi": "CHE - Switzerland",
    "timorleste": "TLS - Timor Leste",
    "tochucdantocthongnhat": "UNO - United Nations Organization",
    "togo": "TGO - Togo",
    "tonga": "TON - Tonga",
    "trungquoc": "CHN - China",
    "trungquocailoan": "CHN - China",
    "tunidi": "TUN - Tunisia",
    "tunisia": "TUN - Tunisia",
    "tuocmenixtan": "TKM - Turkmenistan",
    "turkey": "TUR - Turkey",
    "turkmenistan": "TKM - Turkmenistan",
    "tuvalu": "TUV - Tuvalu",
    "uc": "D - Germany",
    "ucraina": "UKR - Ukraine",
    "udobekixtan": "UZB - Uzbekistan",
    "uganda": "UGA - Uganda",
    "ukraine": "UKR - Ukraine",
    "unitedarabemirates": "ARE - United Arab Emirates",
    "unitedkingdom": "GBD - United Kingdom British Territories Citizen",
    "unitednationsorganization": "UNO - United Nations Organization",
    "unitedstates": "USA - United States of America",
    "urugoay": "URY - Uruguay",
    "uruguay": "URY - Uruguay",
    "uzbekistan": "UZB - Uzbekistan",
    "vanuatu": "VUT - Vanuatu",
    "vaticancity": "VAT - Holy See (Vatican City State )",
    "vaticang": "VAT - Holy See (Vatican City State )",
    "veneduela": "VEN - Venezuela",
    "venezuela": "VEN - Venezuela",
    "vietnam": "VNM - Viet Nam",
    "vungatthuocanhoanoduong": "IOT - British India Ocean Territory",
    "vungthuophap": "FRA - France",
    "vuongquocnauy": "NOR - Norway",
    "westernsamoa": "WSM - Western Samoa",
    "xamoa": "WSM - Western Samoa",
    "xanhluxia": "LCA - Saint Lucia",
    "xanmarino": "SMR - San Marino",
    "xcolent": "SC- - Scotland",
    "xecbia": "SRB - Serbia",
    "xenegan": "SEN - Senegal",
    "xingapo": "SGP - Singapore",
    "xlovakia": "SVK - Slovakia",
    "xoadilen": "SWZ - Swaziland",
    "xomali": "SOM - Somalia",
    "xrilanca": "LKA - Sri Lanka",
    "xuang": "SDN - Sudan",
    "xurinam": "SUR - Suriname",
    "y": "ITA - Italy",
    "yemen": "YEM - Yemen",
    "zambia": "ZMB - Zambia",
    "zimbabwe": "ZWE - Zimbabwe"
}

# ── Sửa các mục tra SAI NƯỚC (nguy hiểm hơn thiếu: vẫn ra "một" kết quả nên
# KHÔNG hiện cảnh báo nào, khách bị khai sai quốc tịch lên hồ sơ công an) ──
#
# "United Kingdom" từng trỏ sang GBD (British Territories Citizen). GBR mới
# là United Kingdom.
NAT_NORM["unitedkingdom"] = "GBR - United Kingdom"
#
# "Úc" từng trỏ sang Đức! Khoá "uc" vốn sinh ra từ "Đức" theo cách chuẩn hoá
# cũ nuốt mất chữ đ ("Đức" → "uc"), rồi chiếm luôn chỗ của "Úc" — khách Úc bị
# khai thành khách Đức. Nay "Đức" → "duc" nên trả "uc" lại đúng cho Úc.
NAT_NORM["uc"] = "AUS - Australia"
NAT_NORM["duc"] = "D - Germany"

# ── Bổ sung tên tiếng Việt thường dùng còn thiếu ──
# Mỗi mục dưới đây CHỈ trỏ tới một giá trị ĐÃ CÓ SẴN trong bảng, không tạo mã
# quốc gia mới. Thiếu chúng thì khách ghi tên nước bằng tiếng Việt thông dụng
# sẽ ra hồ sơ không có mã.
for _vn_alias, _vn_target in (
    ("nga", "RUS - Russia"),            # đã có: lienbangnga, russia
    ("sec", "CZE - Czech Republic"),    # đã có: conghoasec, czechrepublic
    ("dongtimor", "TLS - Timor Leste"), # đã có: ongtimo, timorleste
    ("dailoan", "CHN - China"),         # theo đúng bảng: chinataiwan, trungquocailoan → CHN
):
    assert _vn_target in NAT_NORM.values(), f"mã {_vn_target} không có sẵn trong bảng"
    NAT_NORM[_vn_alias] = _vn_target

# PMS xuất quốc tịch bằng TÊN TIẾNG ANH ĐẦY ĐỦ ("United States of America",
# "Philippines", "Myanmar"...) trong khi bảng trên chỉ có tên Việt và vài
# alias rút gọn ("unitedstates", "philippin") → những khách đó ra hồ sơ KHÔNG
# CÓ MÃ QUỐC TỊCH. Tự sinh khoá từ chính tên hiển thị của mỗi mục, nên mọi
# mục đều tra được bằng đúng tên tiếng Anh của nó. setdefault: không đè khoá
# đã có sẵn, chỉ bù chỗ thiếu.
for _nat_val in list(NAT_NORM.values()):
    if ' - ' in _nat_val:
        NAT_NORM.setdefault(_norm_nat(_nat_val.split(' - ', 1)[1]), _nat_val)

def lookup_nat_kbtt(raw):
    """Khớp thông minh: chuẩn hóa dấu/khoảng trắng để tìm mã quốc tịch."""
    if not raw: return ''
    raw = str(raw).strip()
    # Thử cả 2 cách chuẩn hoá: bản đúng (đ→d) và bản cũ (đ bị nuốt) — các khoá
    # trong NAT_NORM được gõ theo bản cũ nên vẫn phải tra được bằng nó.
    for key in (_norm_nat(raw), _norm_nat_legacy(raw)):
        if key in NAT_NORM:
            return NAT_NORM[key]
    # already in CODE - Name form?
    if _re.match(r'^[A-Z]{2,3} - ', raw):
        return raw
    return raw  # unknown -> keep original (sẽ hiện cảnh báo)

NAT_DK14 = {
    'AFG':'Afganistan  ( Ap-ga-ni-xtan )','ZAF':'Africa (South)  ( Nam Phi )',
    'ALB':'Albania  ( An-ba-ni )','DZA':'Algieria  ( An-giê-ri )',
    'ASM':'American Samoa  ( Đông Sa-moa )','AND':'Andorra  ( Công quốc An-đơ-ra )',
    'ATA':'Antarctica  ( Nam Cực )','AGO':'Angola  ( Ăng-gô-la )',
    'AIA':'Anguilla  ( Ăng-gui-la )','ARG':'Argentina  ( Ac-hen-ti-na )',
    'ARM':'Armenia  ( Ac-mê-ni-a )','ABW':'Aruba  ( A-ru-ba )',
    'AUS':'Australia  ( Ô-xtrây-li-a )','AUT':'Austria  ( áo )',
    'AZE':'Azerbaijan  ( A-déc-bai-gian )','BHS':'Bahamas  ( Ba-ha-ma )',
    'BHR':'Bahrain  ( Ba-ra-in )','BGD':'Bangladesh  ( Băng-la-đét )',
    'BRB':'Barbados  ( Bác-ba-đốt )','BLR':'Belarus  ( Bê-la-rút )',
    'BEL':'Belgium  ( Bỉ )','BLZ':'Belize  ( Bê-li-xê )',
    'BEN':'Benin  ( Bê-nanh )','BMU':'Bermuda  ( Béc-mu-đa )',
    'BTN':'Bhutan  ( Bu-tan )','BOL':'Bolivia  ( Bô-li-vi-a )',
    'BIH':'Bosnia and Herzegovina  ( Bô-xni-a Héc-dê-gô-vi-na )','BWA':'Botswana  ( Bốt-xoa-na )',
    'BRA':'Brazil  ( Bra-din )','BRN':'Brunei Darussalam  ( Đa-ru-xa-lem thuộc Brunei )',
    'BGR':'Bulgaria  ( Bun-ga-ri )','BFA':'Burkina Faso  ( Buốc-ki-na Pha-xô )',
    'BDI':'Burundi  ( Bu-run-đi )','CMR':'Cameroon  ( Ca-mơ-run )',
    'CAN':'Canada  ( Ca-na-da )','COL':'Colombia  ( Cô-lôm-bi-a )',
    'COM':'Comoros  ( Cô-mo )','COG':'Congo  ( Công-gô )',
    'COK':'Cook Islands  ( Quần đảo Cúc )','CRI':'Costa Rica  ( Cô-xta Ri-ca )',
    'HRV':'Croatia  ( Crô-a-ti-a )','CUB':'Cuba  ( Cu Ba )',
    'CYP':'Cyprus  ( Đảo Síp )','CZE':'Czech Republic  ( Cộng hoà Séc )',
    'TCD':'Chad  ( Sát )','CHL':'Chile  ( Chi-lê )',
    'CHN':'China  ( Trung Quốc )','TWN':'China (Taiwan)  ( Trung Quốc (Đài Loan) )',
    'DNK':'Denmark  ( Đan Mạch )','DJI':'Djibouti  ( Đi-bô-u-ti )',
    'DMA':'Dominica  ( Đô-mi-ni-ca )','ECU':'Ecuador  ( Ê-cu-a-đo )',
    'EGY':'Egypt  ( Ai Cập )','SLV':'El Salvador  ( En Xan-va-đo )',
    'GNQ':'Equatorial Guinea  ( Ghi-nê Xích đạo )','ERI':'Eritrea  ( Ê-ri-tơ-ri-a )',
    'EST':'Estonia  ( Ê-xtô-ni-a )','ETH':'Ethiopia  ( Ê-ti-ô-pi-a )',
    'FJI':'Fiji  ( Fi-ji )','FIN':'Finland  ( Phần Lan )',
    'FRA':'France  ( Pháp )','GAB':'Gabon  ( Ga-bông )',
    'GMB':'Gambia  ( Găm-bi-a )','GEO':'Georgia  ( Gru-di-a )',
    'DEU':'Germany  ( CH Liên bang Đức )','GHA':'Ghana  ( Ga-na )',
    'GRC':'Greece  ( Hy Lạp )','GRL':'Greenland  ( Grin-lơn )',
    'GTM':'Guatemala  ( Goa-tê-ma-la )','GIN':'Guinea  ( Ghi-nê )',
    'GNB':'Guinea-Bissau  ( Ghi-nê Bít-xao )','GUY':'Guyana  ( Gui-na )',
    'HTI':'Haiti  ( Ha-i-ti )','HND':'Honduras  ( Hon-du-rat )',
    'HKG':'HongKong  ( Hồng-Kông )','HUN':'Hungari  ( Hung-ga-ri )',
    'ISL':'Iceland  ( Ai-xơ-len )','IND':'India  ( Ân Độ )',
    'IDN':'Indonesia  ( In-đô-nê-xi-a )','IRN':'Iran  ( CH Hồi giáo I-ran )',
    'IRQ':'Iraq  ( I-rắc )','IRL':'Ireland  ( Ai-rơ-len )',
    'ISR':'Israel  ( I-xra-en )','ITA':'Italy  ( I-ta-li-a )',
    'JAM':'Jamaica  ( Ja-mai-ca )','JPN':'Japan  ( Nhật Bản )',
    'JOR':'Jordan  ( Joc-đan )','CAP':'Kampuchea  ( Căm-pu-chia )',
    'KAZ':'Kazakhstan  ( Ka-dắc-xtan )','KEN':'Kenya  ( Kê-ni-a )',
    'KOR':'Korea (South)  ( CH Hàn Quốc )','PRK':'Korea Democratic Peoples Republic  ( CHDCND Triều Tiên )',
    'KWT':'Kuwait  ( Cô-oét )','KGZ':'Kyrgyzstan  ( Kiếc-ghi-di-a )',
    'LAO':'Laos  ( CHDCND Lào )','LVA':'Latvia  ( Lát-vi-a )',
    'LBN':'Lebanon  ( Li-ban )','LSO':'Lesotho  ( Lê-xô-thô )',
    'LBR':'Liberia  ( Li-bê-ri-a )','LBY':'Libya  ( Gia-ma-hi-ri-i-a A-rập Li-bi Nhân dân )',
    'LIE':'Liechtenstein  ( Công quốc Lích-ten-xtên )','LTU':'Lithuania  ( Lit-hua-ni-a )',
    'LUX':'Luxembourg  ( Luých-xem-bua )','MAC':'Macau  ( Ma Cao )',
    'MDG':'Madagascar  ( Ma-đa-ga-xca )','MWI':'Malawi  ( Ma-la-uy )',
    'MYS':'Malaysia  ( Ma-lai-xi-a )','MDV':'Maldives  ( Man-đi-vơ )',
    'MLI':'Mali  ( Ma-li )','MLT':'Malta  ( Man-ta )',
    'MHL':'Marshall Islands  ( Quần đảo Mác-san )','MRT':'Mauritania  ( Mô-ra-ta-ni )',
    'MUS':'Mauritius  ( Mô-ri-xơ )','MEX':'Mexico  ( Mê-xi-cô )',
    'MDA':'Moldova  ( Môn-đô-va )','MCO':'Monaco  ( Công quốc Mô-na-cô )',
    'MNE':'Montenegro  ( Môn-tê-nê-grô )','MNG':'Mongolia  ( Mông Cổ )',
    'MAR':'Morocco  ( Ma-rốc )','MOZ':'Mozambique  ( Mô-dăm-bích )',
    'MMR':'Myanmar (Burma)  ( Mi-an-ma )','NAM':'Namibia  ( Na-mi-bi-a )',
    'NPL':'Nepal  ( Nê-pan )','NLD':'Netherland  ( Hà Lan )',
    'NZL':'New Zealand  ( Niu Di-lân )','NIC':'Nicaragua  ( Ni-ca-ra-goa )',
    'NER':'Niger  ( Ni-giê )','NGA':'Nigeria  ( Ni-giê-ri-a )',
    'NOR':'Norway  ( Vương quốc Na-uy )','OMN':'Oman  ( Ô-man )',
    'PAK':'Pakistan  ( Pa-ki-xtan )','PLW':'Palau  ( Pa-lau )',
    'PSE':'Palestine  ( Pa-le-xtin )','PAN':'Panama  ( Pa-na-ma )',
    'PNG':'Papua New Guinea  ( Pa-pua Niu Ghi-nê )','PRY':'Paraguay  ( Pa-ra-goay )',
    'PER':'Peru  ( Pê-ru )','POL':'Poland  ( Ba Lan )',
    'PRT':'Portugal  ( Bồ Đào Nha )','PRI':'Puerto Rico  ( Pu-éc-tô Ri-cô )',
    'PHL':'Philippine  ( Phi-líp-pin )','QAT':'Qatar  ( Qua-ta )',
    'ROU':'Romania  ( Ru-ma-ni )','RUS':'Russia  (Liên bang Nga)',
    'RWA':'Rwanda  ( Ru-an-đa )','LCA':'Saint Lucia  ( Xanh Lu-xi-a )',
    'SMR':'San Marino  ( Xan Ma-ri-nô )','SAU':'Saudi Arabia  ( A-rập Xau-đi )',
    'GBR':'United Kingdom  ( Liên hiệp Vương quốc Anh và Bắc Ailen )',
    # Một số hệ thống PMS ghi Vương quốc Anh là GBS thay vì GBR — nhận cả hai
    'GBS':'United Kingdom  ( Liên hiệp Vương quốc Anh và Bắc Ailen )',
    'SEN':'Senegal  ( Xe-ne-gan )',
    'SRB':'Serbia  ( Xéc-bi-a )','SYC':'Seychelles  ( Quần đảo Xây-sen )',
    'SGP':'Singapore  ( Xin-ga-po )','SVK':'Slovakia  ( Xlô-va-ki-a )',
    'SVN':'Slovenia  ( Slo-vê-ni-a )','SOM':'Somalia  ( Xô-ma-li )',
    'ESP':'Spain  ( Tây Ban Nha )','LKA':'Srilanka  ( Xri-Lan-ca )',
    'SDN':'Sudan  ( Xu-đăng )','SWE':'Sweden  ( Thuỵ Điển )',
    'CHE':'Switzerland  ( Thuỵ Sĩ )','TJK':'Tajikistan  ( Ta-gi-ki-xtan )',
    'TGO':'Togo  ( Tô-gô )','TON':'Tonga  ( Tôn-ga )',
    'TUN':'Tunisia  ( Tu-ni-di )','TUR':'Turkey  ( Thổ Nhĩ Kỳ )',
    'TKM':'Turkmenistan  ( Tuốc-mê-ni-xtan )','TUV':'Tuvalu  ( Tu-va-lu )',
    'THA':'Thailand  ( Thái Lan )','UGA':'Uganda  ( U-gan-da )',
    'UKR':'Ukraine  ( U-crai-na )','ARE':'United Arab Emirates  ( A-rập thống nhất )',
    'TZA':'United Republic of Tanzania  ( CH thống nhất Tan-da-ni-a )',
    'USA':'United States  ( Mỹ )','URY':'Uruguay  ( U-ru-goay )',
    'UZB':'Uzbekistan  ( U-dơ-bê-ki-xtan )','VUT':'Vanuatu  ( Va-nu-a-tu )',
    'VEN':'Venezuela  ( Vê-nê-du-ê-la )','VNM':'Vietnam  ( Việt Nam )',
    'YEM':'Yemen  ( Y-ê-men )','ZMB':'Zambia  ( Dăm-bi-a )',
    'ZWE':'Zimbabwe  ( Dim-ba-bu-ê )',
}
LOAI_GIAY = {
    'Căn cước công dân':'8 - Thẻ Căn Cước','Hộ chiếu':'4 - Hộ chiếu',
    'Chứng minh nhân dân':'2 - Thẻ CMND','Căn cước':'1 - Thẻ CCCD',
    'Giấy khai sinh':'5 - Giấy khai sinh',
}
TINH = {
    'DAK LAK':'605 - Đắk Lắk','DAK NONG':'607 - Đắk Nông','PHU YEN':'509 - Phú Yên',
    'HO CHI MINH':'701 - TP. Hồ Chí Minh','THP HO CHI MINH':'701 - TP. Hồ Chí Minh',
    'HCM':'701 - TP. Hồ Chí Minh','TP HCM':'701 - TP. Hồ Chí Minh',
    'GIA LAI':'603 - Gia Lai','QUANG NGAI':'505 - Quảng Ngãi','LAM DONG':'703 - Lâm Đồng',
    'LAM DONG.':'703 - Lâm Đồng','KHANH HOA':'511 - Khánh Hòa','BEN TRE':'817 - Bến Tre',
    'VINH LONG':'809 - Vĩnh Long','NINH THUAN':'513 - Ninh Thuận','DONG NAI':'713 - Đồng Nai',
    'TIEN GIANG':'807 - Tiền Giang','LONG AN':'801 - Long An','BINH DUONG':'707 - Bình Dương',
    'BINH THUAN':'705 - Bình Thuận','CAN THO':'815 - TP. Cần Thơ','DA NANG':'501 - TP. Đà Nẵng',
    'HA NOI':'101 - TP. Hà Nội','BA RIA VUNG TAU':'711 - Bà Rịa - Vũng Tàu',
    'DONG THAP':'803 - Đồng Tháp','AN GIANG':'805 - An Giang','KIEN GIANG':'819 - Kiên Giang',
    'HA TINH':'407 - Hà Tĩnh','BINH DINH':'507 - Bình Định','VIET NAM':'',
}


def _norm_addr(s):
    """Chuẩn hóa tên tỉnh/phường để so khớp: bỏ dấu, hoa/thường, bỏ tiền tố
    Xã/Phường/Thị trấn/Đặc khu/Tỉnh/TP..., bỏ ký tự không phải chữ/số.

    Bảng tra (load_vn_admin_lookup) cũng dựng khoá bằng chính hàm này nên cả
    hai phía luôn khớp nhau."""
    s = _strip_accents(s).strip()
    s = s.lower()
    s = _re.sub(r'^(xa|phuong|thi tran|dac khu|tinh|thanh pho|tp\.?)\s+', '', s)
    s = _re.sub(r'[^a-z0-9 ]', ' ', s)
    s = _re.sub(r'\s+', ' ', s).strip()
    return s

@st.cache_resource(show_spinner=False)
def load_vn_admin_lookup():
    """Đọc bảng TINH_THANH + PHUONG_XA nhúng sẵn trong chính mẫu VNM (đúng
    danh mục dropdown DANH_MUC đang dùng, không phải bảng ngoài) → dict tra
    cứu theo tên đã chuẩn hóa, dùng để tự điền mã Tỉnh/Thành + Phường/Xã
    thay vì bỏ trống như trước."""
    wb = load_workbook(io.BytesIO(load_template('vnm')))
    prov_by_norm = {}
    prov_by_code = {}
    for row in wb['TINH_THANH'].iter_rows(min_row=2, values_only=True):
        matt, tentt, display = row[0], row[1], row[2]
        if not matt or not display:
            continue
        prov_by_code[str(matt)] = display
        prov_by_norm.setdefault(_norm_addr(tentt), display)
    ward_by_norm = {}
    for row in wb['PHUONG_XA'].iter_rows(min_row=2, values_only=True):
        ma, ten, matt, display = row[0], row[1], row[2], row[3]
        if not ma or not display:
            continue
        ward_by_norm.setdefault(_norm_addr(ten), []).append((str(matt), display))
    return prov_by_norm, prov_by_code, ward_by_norm

def lookup_province_vnm(raw):
    """Mã Tỉnh/Thành cho cột 'TỈNH/ THÀNH PHỐ' của mẫu VNM — khớp đúng tên
    trong chính mẫu (TINH_THANH) trước, dự phòng bảng TINH cũ (alias/viết
    tắt như 'TP HCM'), giữ nguyên raw nếu không khớp được (không đoán bừa)."""
    raw = str(raw or '').strip()
    if not raw:
        return ''
    prov_by_norm, _, _ = load_vn_admin_lookup()
    hit = prov_by_norm.get(_norm_addr(raw))
    if hit:
        return hit
    legacy = TINH.get(raw.upper())
    if legacy:
        return legacy
    return raw

def lookup_ward_vnm(raw, prov_display=None):
    """Mã Phường/Xã cho cột 'PHƯỜNG/ XÃ/ ĐẶC KHU' của mẫu VNM. Tên phường có
    thể trùng giữa nhiều tỉnh, nên khi đã biết tỉnh (prov_display dạng
    '101 - TP. Hà Nội') sẽ ưu tiên khớp đúng phường thuộc tỉnh đó. Không tự
    suy đoán khi mơ hồ — trả lại raw để lễ tân tự kiểm tra thay vì điền sai.
    Trả về (giá_trị_để_điền, đã_khớp: bool, tỉnh_suy_ra_được: str|None)."""
    raw = str(raw or '').strip()
    if not raw:
        return '', False, None
    _, prov_by_code, ward_by_norm = load_vn_admin_lookup()
    cands = ward_by_norm.get(_norm_addr(raw)) or []
    if not cands:
        return raw, False, None
    prov_code = None
    if prov_display:
        m = _re.match(r'^(\d+)\s*-', str(prov_display))
        if m:
            prov_code = m.group(1)
    if prov_code:
        scoped = [d for c, d in cands if c == prov_code]
        if scoped:
            return scoped[0], True, None
    if len(cands) == 1:
        c, d = cands[0]
        return d, True, (None if prov_display else prov_by_code.get(c))
    return raw, False, None


# Ten duoc module nay so huu — liet ke tuong minh de `import *` lay
# duoc ca helper co gach duoi dau.
__all__ = [
    'LOAI_GIAY', 'NAT_DK14', 'NAT_NORM', 'TINH', '_norm_addr', '_norm_nat',
    '_norm_nat_legacy', 'load_vn_admin_lookup', 'lookup_nat_kbtt',
    'lookup_province_vnm', 'lookup_ward_vnm'
]
