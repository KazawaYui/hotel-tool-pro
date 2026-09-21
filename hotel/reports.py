"""Kiểm tra chất lượng dữ liệu, cảnh báo visa, báo cáo ngày, sổ giao ca."""

import streamlit as st
import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from copy import copy
import xlrd, datetime, io, zipfile, base64, os, json
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.lib.colors import white, black
import unicodedata as _ud, re as _re

from .common import _fmt_room, _gv, _norm_name, _norm_pp, _to_date, now_vn
from .lookups import lookup_nat_kbtt

# ── Tiện ích nghiệp vụ lễ tân (tỷ giá, kiểm tra dữ liệu, visa, báo cáo, giao ca) ──
@st.cache_data(ttl=600, show_spinner=False)
def fetch_vcb_rates():
    """Lấy tỷ giá CHUYỂN KHOẢN USD/EUR → VND từ Vietcombank (cache 10 phút).
    Chỉ là tiện ích — lễ tân vẫn nhập tay được nếu mạng/VCB lỗi."""
    import urllib.request
    import xml.etree.ElementTree as ET
    url = 'https://portal.vietcombank.com.vn/Usercontrols/TVPortal.TyGia/pXML.aspx'
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req, timeout=10) as resp:
        root = ET.fromstring(resp.read())
    rates = {}
    for ex in root.iter('Exrate'):
        code = (ex.get('CurrencyCode') or '').upper()
        if code in ('USD', 'EUR'):
            try:
                rates[code] = float((ex.get('Transfer') or '').replace(',', ''))
            except ValueError:
                pass
    if 'USD' not in rates:
        raise ValueError('không đọc được tỷ giá USD trong dữ liệu VCB')
    return rates, now_vn().strftime('%H:%M %d/%m/%Y')


def validate_guests(df):
    """Kiểm tra chất lượng dữ liệu khách TRƯỚC khi nộp hồ sơ KBTT/VNM/ĐK14.
    🔴 = lỗi dễ khiến công an trả hồ sơ · 🟡 = nên kiểm tra lại trước khi nộp."""
    issues = []
    for idx, row in df.iterrows():
        line = idx + 2  # số dòng trên Excel gốc (dòng 1 là header)
        ht = str(_gv(row, 'HỌ TÊN ', 'HỌ TÊN') or '').strip()
        sp = _fmt_room(_gv(row, 'SỐ PHÒNG'))
        intl = str(_gv(row, 'LOẠI KHÁCH') or '').strip() == 'Quốc tế'
        def add(sev, msg):
            issues.append({'Mức độ': sev, 'Dòng': line, 'Họ tên': ht or '(trống)',
                           'Phòng': sp, 'Vấn đề': msg})
        if not ht:
            add('🔴', 'Thiếu họ tên')
        sg = str(_gv(row, 'SỐ GIẤY TỜ') or '').strip()
        if not sg:
            add('🔴' if intl else '🟡',
                'Thiếu số giấy tờ' + (' (hộ chiếu bắt buộc cho KBTT)' if intl else ''))
        elif intl and not _re.fullmatch(r'[A-Za-z0-9]{4,15}', sg.replace(' ', '')):
            add('🟡', f'Số hộ chiếu có ký tự lạ: "{sg}"')
        if _gv(row, 'NGÀY SINH') is None:
            add('🟡', 'Thiếu ngày sinh')
        if str(_gv(row, 'GIỚI TÍNH') or '').strip() not in ('Nam', 'Nữ'):
            add('🟡', 'Giới tính trống/không chuẩn (cần "Nam" hoặc "Nữ")')
        if not sp:
            add('🟡', 'Thiếu số phòng')
        nd = _to_date(_gv(row, 'NGÀY ĐẾN'))
        ni = _to_date(_gv(row, 'NGÀY ÐI', 'NGÀY ĐI'))
        if nd and ni and ni < nd:
            add('🔴', f'Ngày đi {ni.strftime("%d/%m/%Y")} TRƯỚC ngày đến {nd.strftime("%d/%m/%Y")}')
        if intl:
            qt = str(_gv(row, 'QUỐC TỊCH') or '').strip()
            if not qt:
                add('🔴', 'Thiếu quốc tịch')
            elif not _re.match(r'^[A-Z]{2,3} - ', str(lookup_nat_kbtt(qt))):
                add('🟡', f'Quốc tịch chưa có mã: "{qt}"')
    return pd.DataFrame(issues, columns=['Mức độ', 'Dòng', 'Họ tên', 'Phòng', 'Vấn đề'])

def check_visa_expiry(df_intl, visa_map=None):
    """Gom hạn tạm trú/visa từng khách quốc tế — CHỈ tin dữ liệu từ file Visa
    rời do lễ tân chủ động upload (khớp theo hộ chiếu trước, tên là dự phòng).
    KHÔNG tự đọc cột 'TẠM TRÚ'/'TẠM TRÚ ĐẾN NGÀY' có sẵn trong file dữ liệu
    khách (customer.xls) nữa — cột đó thường do PMS tự điền mặc định, không
    phải visa đã xác nhận thật, tin nhầm có thể điền sai lên hồ sơ khai báo
    công an. → list dict ngày dạng ISO, lưu được vào session để lọc lại theo
    số ngày cảnh báo mà không cần xử lý lại."""
    visa_map = visa_map or {}
    by_pp = visa_map.get('by_pp', {}) if isinstance(visa_map, dict) else {}
    by_name = visa_map.get('by_name', {}) if isinstance(visa_map, dict) else {}
    out = []
    for _, row in df_intl.iterrows():
        ht = str(_gv(row, 'HỌ TÊN ', 'HỌ TÊN') or '').strip()
        sg = str(_gv(row, 'SỐ GIẤY TỜ') or '').strip()
        vd_raw = None
        if by_pp or by_name:
            vd_raw = by_pp.get(_norm_pp(sg)) or by_name.get(_norm_name(ht))
        vd = _to_date(vd_raw)
        if vd is None:
            continue
        ni = _to_date(_gv(row, 'NGÀY ÐI', 'NGÀY ĐI'))
        out.append({'name': ht, 'room': _fmt_room(_gv(row, 'SỐ PHÒNG')),
                    'nat': str(_gv(row, 'QUỐC TỊCH') or '').strip(),
                    'visa': vd.isoformat(), 'dep': ni.isoformat() if ni else None})
    return out

def build_handover_xlsx(info, entries):
    """Xuất sổ giao ca thành file Excel in được, có cột 'Đã xử lý' để ca sau tick tay."""
    wb = Workbook(); ws = wb.active; ws.title = 'Giao ca'
    thin = Side(style='thin'); bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, w in zip('ABCDEF', [5, 8, 26, 9, 58, 10]):
        ws.column_dimensions[col].width = w
    ws.merge_cells('A1:F1')
    c = ws.cell(1, 1)
    c.value = (f"SỔ GIAO CA — {info.get('date', '')} — {info.get('shift', '')}"
               f" — Lễ tân: {info.get('staff', '') or '…'}")
    c.font = Font(name='Times New Roman', size=14, bold=True)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30
    for ci, h in enumerate(['STT', 'Giờ ghi', 'Phân loại', 'Phòng', 'Nội dung bàn giao', 'Đã xử lý'], 1):
        cell = ws.cell(2, ci); cell.value = h
        cell.font = Font(name='Times New Roman', size=11, bold=True)
        cell.fill = PatternFill('solid', fgColor='DDEBF7'); cell.border = bdr
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, e in enumerate(entries, 1):
        ws.row_dimensions[i + 2].height = 32
        vals = [i, e.get('time', ''), e.get('cat', ''), e.get('room', ''), e.get('note', ''), '☐']
        for ci, v in enumerate(vals, 1):
            cell = ws.cell(i + 2, ci); cell.value = v
            cell.font = Font(name='Times New Roman', size=11); cell.border = bdr
            cell.alignment = Alignment(horizontal='left' if ci == 5 else 'center',
                                       vertical='center', wrap_text=True)
    return wb

def build_daily_report(date_str, daily, arr_stats, recon, reconr):
    """Báo cáo ngày 1 trang (Excel) cho quản lý — tổng hợp mọi số liệu các công cụ
    đã chạy trong phiên; phần nào chưa chạy thì tự bỏ qua."""
    wb = Workbook(); ws = wb.active; ws.title = 'Bao cao ngay'
    thin = Side(style='thin'); bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, w in zip('AB', [40, 22]):
        ws.column_dimensions[col].width = w
    bold = Font(name='Times New Roman', size=11, bold=True)
    norm = Font(name='Times New Roman', size=11)
    fill_h = PatternFill('solid', fgColor='DDEBF7')
    ws.merge_cells('A1:B1')
    t = ws.cell(1, 1); t.value = f"BÁO CÁO NGÀY {date_str} — TÂN HOTEL (FRONT OFFICE)"
    t.font = Font(name='Times New Roman', size=14, bold=True)
    t.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28
    r = 3
    def section(title):
        nonlocal r
        ws.merge_cells(f'A{r}:B{r}')
        c = ws.cell(r, 1); c.value = title; c.font = bold; c.fill = fill_h; c.border = bdr
        ws.cell(r, 2).border = bdr
        r += 1
    def kv(label, value):
        nonlocal r
        a = ws.cell(r, 1); a.value = label; a.font = norm; a.border = bdr
        b = ws.cell(r, 2); b.value = value; b.font = bold; b.border = bdr
        b.alignment = Alignment(horizontal='center')
        r += 1
    if daily:
        section('1. KHÁCH LƯU TRÚ (file dữ liệu khách)')
        kv('Tổng khách', daily.get('total'))
        kv('Khách quốc tế', daily.get('intl'))
        kv('Khách Việt Nam', daily.get('vn'))
        if daily.get('rooms_cnt'):
            kv('Số phòng có khách', daily.get('rooms_cnt'))
        if daily.get('avg_nights'):
            kv('Số đêm lưu trú bình quân', daily.get('avg_nights'))
        kv('Trẻ em (GKS) + Giấy bảo lãnh (GBL)', f"{daily.get('gks', 0)} + {daily.get('gbl', 0)}")
        r += 1
        if daily.get('nat_top'):
            section('2. TOP QUỐC TỊCH')
            for nat, cnt in daily['nat_top']:
                kv(nat, cnt)
            r += 1
    if arr_stats:
        section('3. BOOKING ĐẾN & THANH TOÁN (file ARR)')
        kv('Số booking arrival', arr_stats.get('bookings'))
        kv('Số phòng arrival', arr_stats.get('rooms'))
        if arr_stats.get('ota') is not None:
            kv('Booking qua OTA', arr_stats.get('ota'))
        kv('Cần cà thẻ (CÀ THẺ)', arr_stats.get('ca_the'))
        kv('Cần thu tiền (THU TIỀN)', arr_stats.get('thu_tien'))
        kv('Xem lại BU', arr_stats.get('xem_lai_bu'))
        kv('FOC Late C/O', arr_stats.get('foc_lco'))
        r += 1
    if recon:
        section('4. ĐỐI CHIẾU LƯU TRÚ NGƯỜI NƯỚC NGOÀI')
        kv('Khách chưa đăng ký lưu trú', len(recon.get('chua_dk', [])))
        kv('Có trên lưu trú, thiếu trên Smile', len(recon.get('thua', [])))
        kv('Đăng ký trùng', len(recon.get('dup', [])))
        r += 1
    if reconr:
        section('5. ĐỐI CHIẾU HỆ THỐNG PHÒNG')
        kv('Phòng chưa đăng ký', len(reconr.get('room_chua', [])))
        kv('Phòng thừa trong file', len(reconr.get('room_thua', [])))
        kv('Phòng trùng trong file', len(reconr.get('sys_dup', [])))
        r += 1
    r += 1
    f = ws.cell(r, 1)
    f.value = f"Xuất lúc {now_vn().strftime('%H:%M %d/%m/%Y')} — Hotel Tool Pro"
    f.font = Font(name='Times New Roman', size=9, italic=True, color='FF888888')
    return wb


# Ten duoc module nay so huu — liet ke tuong minh de `import *` lay
# duoc ca helper co gach duoi dau.
__all__ = [
    'build_daily_report', 'build_handover_xlsx', 'check_visa_expiry', 'fetch_vcb_rates',
    'validate_guests'
]
