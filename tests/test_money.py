"""Quy đổi tỷ giá + chia giá phòng connecting.

Đây là vùng đã hai lần sinh lỗi thật (PR #26/#27 và #28), mỗi lỗi đều sai
tiền triệu trên hồ sơ. Test ở đây khoá chặt hành vi đúng.
"""
import io

import pytest
from openpyxl import Workbook, load_workbook


RATE = 29860.07


def _qllt(rows, headers=('MÃ CHECKIN', 'SỐ PHÒNG', 'ĐƠN GIÁ')):
    """Dựng file QLLT tối giản trong bộ nhớ."""
    wb = Workbook()
    ws = wb.active
    ws.append(list(headers))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _doc_gia(xlsx_bytes, phong=None):
    ws = load_workbook(io.BytesIO(xlsx_bytes)).active
    h = [c.value for c in ws[1]]
    ci_room, ci_gia = h.index('SỐ PHÒNG') + 1, h.index('ĐƠN GIÁ') + 1
    out = []
    for r in range(2, ws.max_row + 1):
        room = str(ws.cell(r, ci_room).value)
        if phong is None or room in phong:
            out.append((room, ws.cell(r, ci_gia).value))
    return out


# ── Danh sách cặp connecting ──────────────────────────────────────────────
def test_du_78_cap_13_tang(app):
    """Nguồn C_p_CNT.xlsx có 13 cột tầng (5→12A, 14, 15, 16, 17) × 6 cặp.
    Bản đầu chỉ chép 9 cột, bản sau 11 cột — đều thiếu."""
    assert len(app.CONNECTING_ROOM_PAIRS) == 78
    assert len(set(app.CONNECTING_ROOM_PAIRS)) == 78, "có cặp bị trùng"


@pytest.mark.parametrize('cap', [
    ('538', '540'), ('742', '744'), ('849', '850'),      # tầng 5, 7, 8
    ('1138', '1140'), ('1241', '1242'),                   # tầng 11, 12
    ('12A30', '12A32'), ('12A41', '12A42'),               # tầng 12A
    ('1428', '1430'), ('1440', '1441'),                   # tầng 14
    ('1528', '1530'), ('1537', '1539'),                   # tầng 15
    ('1628', '1630'), ('1637', '1639'),                   # tầng 16
    ('1728', '1730'), ('1737', '1739'),                   # tầng 17
])
def test_cap_co_trong_danh_sach(app, cap):
    assert cap in app.CONNECTING_ROOM_PAIRS


def test_danh_sach_viet_hoa(app):
    """PMS xuất chữ thường ('12a30'), sơ đồ tầng ghi hoa — danh sách phải ở
    dạng HOA để so khớp sau khi chuẩn hoá."""
    for a, b in app.CONNECTING_ROOM_PAIRS:
        assert a == a.upper() and b == b.upper()


# ── Chia giá ──────────────────────────────────────────────────────────────
def test_chia_doi_tong_le_khong_lam_mat_tien(app):
    """Tổng lẻ: phòng số nhỏ làm tròn xuống, số lớn làm tròn lên — cộng lại
    phải ĐÚNG BẰNG tổng gốc, không được rơi mất 1 đồng."""
    out, rep = app.split_connecting_room_prices(
        _qllt([('B1', '742', 8313043), ('B1', '744', 0)]))
    gia = dict(_doc_gia(out, {'742', '744'}))
    assert gia['742'] == 4156521
    assert gia['744'] == 4156522
    assert gia['742'] + gia['744'] == 8313043
    assert len(rep) == 1


def test_chia_doi_tong_chan(app):
    out, _ = app.split_connecting_room_prices(
        _qllt([('B1', '538', 2000000), ('B1', '540', 0)]))
    gia = dict(_doc_gia(out, {'538', '540'}))
    assert gia['538'] == 1000000 and gia['540'] == 1000000


def test_phong_chu_thuong_van_khop(app):
    out, rep = app.split_connecting_room_prices(
        _qllt([('B1', '12a30', 777777), ('B1', '12a32', 0)]))
    assert len(rep) == 1
    gia = dict(_doc_gia(out, {'12a30', '12a32'}))
    assert gia['12a30'] + gia['12a32'] == 777777


def test_ca_hai_phong_deu_co_gia_thi_khong_dung(app):
    """Chỉ chia khi ĐÚNG MỘT phòng có giá. Cả hai có giá = PMS đã tách sẵn."""
    out, rep = app.split_connecting_room_prices(
        _qllt([('B1', '1138', 3613068), ('B1', '1140', 3642929)]))
    assert rep == []
    assert dict(_doc_gia(out, {'1138', '1140'})) == {'1138': 3613068, '1140': 3642929}


def test_ca_hai_phong_deu_0_thi_khong_dung(app):
    out, rep = app.split_connecting_room_prices(
        _qllt([('B1', '742', 0), ('B1', '744', 0)]))
    assert rep == []


def test_khac_ma_checkin_thi_khong_dung(app):
    """Hai phòng connecting nhưng KHÁC lượt đặt — không được gộp tiền."""
    out, rep = app.split_connecting_room_prices(
        _qllt([('B1', '742', 5000000), ('B2', '744', 0)]))
    assert rep == []
    assert dict(_doc_gia(out, {'742', '744'})) == {'742': 5000000, '744': 0}


def test_cap_ngoai_danh_sach_giu_nguyen(app):
    """1604/1606 cùng tầng 16 nhưng KHÔNG phải cặp connecting — không đoán."""
    out, rep = app.split_connecting_room_prices(
        _qllt([('B1', '1604', 5000000), ('B1', '1606', 0)]))
    assert rep == []
    assert dict(_doc_gia(out, {'1604', '1606'})) == {'1604': 5000000, '1606': 0}


def test_thieu_cot_thi_tra_nguyen_file(app):
    raw = _qllt([('B1', '742', 100)], headers=('MÃ CHECKIN', 'SỐ PHÒNG', 'GIÁ'))
    out, rep = app.split_connecting_room_prices(raw)
    assert out == raw and rep == []


# ── Quy đổi tỷ giá ────────────────────────────────────────────────────────
def test_quy_doi_ngoai_te_nho(app):
    wb, conv = app.process_xlsx(_qllt([('B1', '742', 278.40)]), RATE)
    ws = wb['customer']
    h = [c.value for c in ws[1]]
    assert ws.cell(2, h.index('ĐƠN GIÁ') + 1).value == round(278.40 * RATE)
    assert conv == 1


def test_khong_quy_doi_gia_da_la_vnd(app):
    """Giá ≥ 1000 coi như đã là VNĐ — không nhân tỷ giá lần nữa."""
    wb, conv = app.process_xlsx(_qllt([('B1', '742', 4156521)]), RATE)
    ws = wb['customer']
    h = [c.value for c in ws[1]]
    assert ws.cell(2, h.index('ĐƠN GIÁ') + 1).value == 4156521
    assert conv == 0


# ── Hồi quy lỗi đã từng xảy ra thật ───────────────────────────────────────
def test_hoi_quy_thu_tu_quy_doi_roi_moi_chia(app):
    """Lỗi thật (PR #27): chia TRƯỚC khi quy đổi thì 278.40 EUR bị int() cắt
    còn 278, chia 139/139, quy đổi từng nửa → 4.150.550 mỗi phòng, tổng
    8.301.100 — lệch 11.943đ so với PMS. Phải quy đổi TRƯỚC, chia SAU."""
    raw = _qllt([('B1', '742', 278.40), ('B1', '744', 0)])

    wb, _ = app.process_xlsx(raw, RATE)              # 1. quy đổi
    out, _ = app.split_connecting_room_prices(app.wb_to_bytes(wb))   # 2. chia
    ws = load_workbook(io.BytesIO(out))['customer']
    h = [c.value for c in ws[1]]
    ci_room, ci_gia = h.index('SỐ PHÒNG') + 1, h.index('ĐƠN GIÁ') + 1
    tong = sum(ws.cell(r, ci_gia).value or 0 for r in range(2, ws.max_row + 1)
               if str(ws.cell(r, ci_room).value) in ('742', '744'))
    assert tong == 8313043, "phải khớp đúng số dư trên PMS"
    assert tong != 8301100, "đây là con số SAI của thứ tự cũ"


# ── Định dạng file xuất ra ────────────────────────────────────────────────
def test_o_du_lieu_duoc_ke_dinh_dang_theo_dong_mau(app):
    """Style được chụp MỘT LẦN từ dòng mẫu rồi dán cho mọi ô (nhanh gấp ~2,5
    lần so với copy từng ô). Test này chặn trường hợp tối ưu quá tay làm mất
    luôn định dạng — file nộp công an sẽ sai mẫu."""
    wb, _ = app.process_xlsx(_qllt([('B1', '742', 278.40),
                                    ('B1', '744', 1500000)]), RATE)
    ws = wb['customer']
    for r in (2, 3):
        c = ws.cell(r, 2)
        assert c.font is not None and c.font.name, f'dòng {r} mất font'
        assert c.border is not None, f'dòng {r} mất viền'
        assert c.alignment is not None, f'dòng {r} mất căn lề'


def test_moi_dong_co_cung_dinh_dang(app):
    """Mọi dòng khách phải trông giống nhau — chụp style dùng chung không
    được làm dòng sau khác dòng trước."""
    rows = [('B1', str(600 + i), 1500000) for i in range(10)]
    wb, _ = app.process_xlsx(_qllt(rows), RATE)
    ws = wb['customer']
    mau = [(str(ws.cell(2, c).font), str(ws.cell(2, c).border),
            ws.cell(2, c).number_format) for c in range(1, ws.max_column + 1)]
    for r in range(3, 12):
        dong = [(str(ws.cell(r, c).font), str(ws.cell(r, c).border),
                 ws.cell(r, c).number_format) for c in range(1, ws.max_column + 1)]
        assert dong == mau, f'dòng {r} khác định dạng với dòng 2'
